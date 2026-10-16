# -*- coding: utf-8 -*-
"""课件 PPTX → 按页 markdown（页号锚点 ## pNN，含表格与备注页）。

用法:
    python pptx2md.py <课件.pptx 或 目录> [输出目录]
目录模式下，目录内所有 .pptx 各生成一份 md；输出目录默认 ./课件文本/。
"""
import sys
from pathlib import Path

from pptx import Presentation

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def shape_lines(shape, depth: int = 0) -> list[str]:
    """抽取一个形状里的文字；表格按行拼、组合形状递归。"""
    out: list[str] = []
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            out += shape_lines(sub, depth + 1)
        return out
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
        return out
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            out.append(text)
    return out


def convert(src: Path, out: Path) -> int:
    prs = Presentation(str(src))
    lines = [
        f"# {src.stem} · 逐页文本",
        "",
        f"> 源：`{src.name}`（{len(prs.slides)} 页）",
        "> 页号＝PPT 页码（第 N 张幻灯片即 pN）。",
        "> 图片、图表里的文字抽不出来，讲课时以你屏上的课件为准。",
        "",
    ]
    for i, slide in enumerate(prs.slides, start=1):
        lines += [f"## p{i}", ""]
        body: list[str] = []
        for shape in slide.shapes:
            body += shape_lines(shape)
        lines += ["\n\n".join(body).strip() or "（本页无文字）", ""]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines += [f"**备注**：{notes}", ""]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"{src.name}: {len(prs.slides)} 页 → {out}")
    return len(prs.slides)


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(1)
    src = Path(sys.argv[1])
    if src.is_dir():
        out_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("课件文本")
        for f in sorted(src.glob("*.pptx")):
            convert(f, out_dir / f"{f.stem}.md")
    else:
        out = Path(sys.argv[2]) / f"{src.stem}.md" if len(sys.argv) > 2 else src.with_suffix(".md")
        convert(src, out)


if __name__ == "__main__":
    main()
